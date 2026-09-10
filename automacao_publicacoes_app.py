import glob
import os
import re
import openpyxl
from openpyxl.styles import PatternFill

NOME_NOVA_ABA = "Publicações - Processadas"
LIMITE_SIMILARIDADE = 0.85

# Cores ANSI
VERMELHO_THONON = "\033[38;2;166;44;43m"  # Tom vinho/terracota da marca
CINZA_THONON = "\033[38;2;110;110;110m"  # Cinza dos blocos inferiores
RESET = "\033[0m"
NEGRITO = "\033[1m"


def exibir_banner():
  """Habilita as cores no terminal do Windows e exibe a logo perfeitamente centralizada."""
  os.system("")  # Ativa suporte ANSI no CMD/PowerShell

  logo = f"""
{VERMELHO_THONON}                  ██████████████████████████████
                  ██████████████████████████████{RESET}

{CINZA_THONON}                  ████████████      ████████████
                  ████████████      ████████████
                  ████████████      ████████████{RESET}

                           {VERMELHO_THONON}{NEGRITO}T H O N O N{RESET}
                        {CINZA_THONON}A D V O G A D O S{RESET}
  """
  print(logo)
  print(
      f"{CINZA_THONON}=================================================================={RESET}"
  )
  print(
      f"          {VERMELHO_THONON}{NEGRITO}SISTEMA DE AUTOMAÇÃO DE PUBLICAÇÕES UNIFICADAS{RESET}"
  )
  print(
      f"{CINZA_THONON}=================================================================={RESET}\n"
  )


def extrair_corpo_texto(texto):
  if not isinstance(texto, str):
    return ""
  if "CONTEUDO:" in texto:
    texto = texto.split("CONTEUDO:", 1)[1]
  elif "TEXTO:" in texto:
    texto = texto.split("TEXTO:", 1)[1]

  texto = re.sub(r"https?://\S+", "", texto)
  texto = re.sub(r"[^\w\s]", " ", texto)
  texto = re.sub(r"\s+", " ", texto)
  return texto.strip().lower()


def calcular_similaridade_palavras(str1, str2):
  s1 = set(str1.split())
  s2 = set(str2.split())
  if not s1 or not s2:
    return 0.0
  return len(s1.intersection(s2)) / max(len(s1), len(s2))


def extrair_numero_processo(texto):
  match = re.search(r"\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}", str(texto))
  return match.group(0) if match else ""


def identificar_area(texto_ou_processo):
  match = re.search(
      r"\d{7}-\d{2}\.\d{4}\.(\d)\.\d{2}\.\d{4}", str(texto_ou_processo)
  )
  if match:
    return "Trabalhista" if match.group(1) == "5" else "Cível"
  return ""

# Define a prioridade dos status - Sucessos (Lançado/Duplicado) têm prioridade
def obter_peso_status(status):
  status_upper = str(status).upper()
  if (
      "ANDAMENTO/TAREFA CRIADA" in status_upper
      or "CONSIDERADO COMO DUPLICADO" in status_upper
  ):
    return 3
  if "PROCESSO ARQUIVADO" in status_upper:
    return 2
  if "PROCESSO NÃO ENCONTRADO" in status_upper:
    return 1
  return 0

def processar_arquivo(caminho_arquivo):
  print(
      f"\n{CINZA_THONON}-> Processando:{RESET} {NEGRITO}{os.path.basename(caminho_arquivo)}{RESET}"
  )
  try:
    wb = openpyxl.load_workbook(caminho_arquivo)
  except Exception as e:
    print(
        f"{VERMELHO_THONON}Erro ao abrir o arquivo (pode estar aberto no Excel): {e}{RESET}"
    )
    return

  nome_pub = next((s for s in wb.sheetnames if "publica" in s.lower()), None)
  nome_rpa = next(
      (
          s
          for s in wb.sheetnames
          if "rpa" in s.lower()
          or "tarefa" in s.lower()
          or "análise" in s.lower()
          or "analise" in s.lower()
      ),
      None,
  )

  if not nome_pub or not nome_rpa:
    print(
        f"{VERMELHO_THONON}Abas necessárias não encontradas no arquivo: {wb.sheetnames}{RESET}"
    )
    return

  ws_pub_orig = wb[nome_pub]
  ws_rpa = wb[nome_rpa]

  if NOME_NOVA_ABA in wb.sheetnames:
    del wb[NOME_NOVA_ABA]

  ws_proc = wb.copy_worksheet(ws_pub_orig)
  ws_proc.title = NOME_NOVA_ABA

  # Mapear RPA
  rpa_headers = {
      str(ws_rpa.cell(row=1, column=c).value).strip().upper(): c
      for c in range(1, ws_rpa.max_column + 1)
  }
  col_rpa_proc = rpa_headers.get("NUM. PROCESSO", 3)
  col_rpa_texto = rpa_headers.get("TEXTO PUBLICAÇÃO", 6)
  col_rpa_status = rpa_headers.get("STATUS", 8)

  rpa_dados = []
  for r in range(2, ws_rpa.max_row + 1):
    texto_raw = ws_rpa.cell(row=r, column=col_rpa_texto).value or ""
    status = str(
        ws_rpa.cell(row=r, column=col_rpa_status).value or ""
    ).strip().upper()
    proc = str(ws_rpa.cell(row=r, column=col_rpa_proc).value or "").strip()

    rpa_dados.append({
        "processo": proc,
        "texto_limpo": extrair_corpo_texto(texto_raw),
        "status": status,
    })

  # Mapear Publicações pelos cabeçalhos (com fallback para o layout atual,
  # caso alguma coluna seja renomeada ou o cabeçalho não seja encontrado)
  pub_headers = {
      str(ws_proc.cell(row=1, column=c).value).strip().upper(): c
      for c in range(1, ws_proc.max_column + 1)
  }
  col_pub_intimacao = pub_headers.get("INTIMAÇÃO", 2)
  col_pub_proc = pub_headers.get("PROCESSO", 4)
  col_pub_classificacao = pub_headers.get("CLASSIFICAÇÃO", 5)
  col_pub_area = pub_headers.get("ÁREA", 6)
  col_pub_diagnostico = pub_headers.get("DIAGNÓSTICO", 7)

  headers_esperados = ["INTIMAÇÃO", "PROCESSO", "CLASSIFICAÇÃO", "ÁREA", "DIAGNÓSTICO"]
  headers_faltando = [h for h in headers_esperados if h not in pub_headers]
  if headers_faltando:
    print(
        f"{VERMELHO_THONON}Aviso: cabeçalho(s) não encontrado(s) em "
        f"'{nome_pub}': {headers_faltando} — usando posição padrão, "
        f"CONFIRA o resultado.{RESET}"
    )

  fill_duplicado = PatternFill(
      start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
  )

  textos_processados_pub = []

  for row in range(2, ws_proc.max_row + 1):
    intimacao_raw = ws_proc.cell(row=row, column=col_pub_intimacao).value or ""
    proc_pub = str(
        ws_proc.cell(row=row, column=col_pub_proc).value or ""
    ).strip()

    if not proc_pub:
      proc_pub = extrair_numero_processo(intimacao_raw)
      ws_proc.cell(row=row, column=col_pub_proc).value = proc_pub

    # Preenche Área
    area_identificada = identificar_area(
        proc_pub if proc_pub else intimacao_raw
    )
    if area_identificada:
      ws_proc.cell(row=row, column=col_pub_area).value = area_identificada

    corpo_pub = extrair_corpo_texto(intimacao_raw)
    if not corpo_pub:
      continue

    # Duplicidade
    if any(
        calcular_similaridade_palavras(corpo_pub, ant) >= LIMITE_SIMILARIDADE
        for ant in textos_processados_pub
    ):
      for col in range(1, ws_proc.max_column + 1):
        ws_proc.cell(row=row, column=col).fill = fill_duplicado
    else:
      textos_processados_pub.append(corpo_pub)

   # Confronto RPA - Busca pelo processo com prioridade de status e similaridade
    melhor_match = None
    melhor_prioridade = -1
    maior_score = -1.0

    for rpa in rpa_dados:
      mesmo_processo = proc_pub and (proc_pub == rpa["processo"])
      sim = calcular_similaridade_palavras(corpo_pub, rpa["texto_limpo"])
      prioridade = obter_peso_status(rpa["status"])

      # Prioriza status de sucesso, em caso de mesmo status, desempata pela similaridade e aceita registro mais recente
      if mesmo_processo:
        if prioridade > melhor_prioridade or (
            prioridade == melhor_prioridade and sim >= maior_score
        ):
          melhor_prioridade = prioridade
          maior_score = sim
          melhor_match = rpa

      elif not proc_pub:
        # Contingência para publicação sem número de processo (apenas texto)
        if sim >= LIMITE_SIMILARIDADE and (
            prioridade > melhor_prioridade
            or (prioridade == melhor_prioridade and sim >= maior_score)
        ):
          melhor_prioridade = prioridade
          maior_score = sim
          melhor_match = rpa

    # Classificação e Diagnóstico com base no melhor registro exato
    if melhor_match:
      status = melhor_match["status"]

      if (
          "ANDAMENTO/TAREFA CRIADA" in status
          or "CONSIDERADO COMO DUPLICADO" in status
      ):
        ws_proc.cell(row=row, column=col_pub_classificacao).value = "Lançado"
        ws_proc.cell(row=row, column=col_pub_diagnostico).value = "OK"
      elif "PROCESSO ARQUIVADO" in status:
        ws_proc.cell(row=row, column=col_pub_classificacao).value = "Não Lançado"
        ws_proc.cell(row=row, column=col_pub_diagnostico).value = (
            "Processo Arquivado"
        )
      elif "PROCESSO NÃO ENCONTRADO" in status:
        ws_proc.cell(row=row, column=col_pub_classificacao).value = "Não Lançado"
        ws_proc.cell(row=row, column=col_pub_diagnostico).value = (
            "Processo não encontrado - verificar"
        )
      else:
        ws_proc.cell(row=row, column=col_pub_classificacao).value = "Não Lançado"
        ws_proc.cell(row=row, column=col_pub_diagnostico).value = (
            f"Status RPA não mapeado: {status}"
        )
    else:
      ws_proc.cell(row=row, column=col_pub_classificacao).value = "Não Lançado"
      ws_proc.cell(row=row, column=col_pub_diagnostico).value = (
          "Não Localizado no RPA"
      )

  wb.save(caminho_arquivo)
  print(f"   {CINZA_THONON}Concluído na aba '{NOME_NOVA_ABA}'.{RESET}")


def main():
  exibir_banner()

  todos_arquivos = glob.glob("*.xlsx")
  arquivos_alvo = [
      f
      for f in todos_arquivos
      if re.search(r"publica[cç][oõ]es unificadas", f, re.IGNORECASE)
      and not f.startswith("~$")
  ]

  if not arquivos_alvo:
    print(
        f"{VERMELHO_THONON}Nenhum arquivo de 'Publicações Unificadas' encontrado nesta pasta.{RESET}"
    )
  else:
    print(
        f"{NEGRITO}{len(arquivos_alvo)} arquivo(s) encontrado(s) para processar.{RESET}"
    )
    for arq in arquivos_alvo:
      processar_arquivo(arq)

  print(
      f"\n{CINZA_THONON}=============================================================={RESET}"
  )
  input(
      f"{NEGRITO}Processamento finalizado. Pressione ENTER para sair...{RESET}"
  )


if __name__ == "__main__":
  main()